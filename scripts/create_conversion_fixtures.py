"""Generate non-sensitive office fixtures for the real converter integration test."""
from pathlib import Path
import sys
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

target = Path(sys.argv[1])
target.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_heading("PrintR conversion check", 0)
document.add_paragraph("Local document support. Test fixture, no personal information.")
document.save(target / "document.docx")
workbook = Workbook()
sheet = workbook.active
sheet.append(["PrintR test", "Total"])
sheet.append(["Sample", 42])
workbook.save(target / "spreadsheet.xlsx")
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[0])
slide.shapes.title.text = "PrintR conversion check"
slide.placeholders[1].text = "Presentation fixture"
deck.save(target / "presentation.pptx")
(target / "document.rtf").write_text(r"{\rtf1\ansi PrintR conversion check.}", encoding="ascii")
print("Created four Office fixtures. Convert the OOXML files to ODT/ODS/ODP with LibreOffice for the remaining tests.")
